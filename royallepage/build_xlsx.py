"""Step 4: assemble the per-province CSVs into the final deliverable workbook.

One sheet per province, an "All" sheet with every agent, and a "Summary"
sheet with agent counts per city (the 45 selected cities across the three
provinces) so the 45-city breakdown can be sanity-checked at a glance.
"""

from __future__ import annotations

import csv
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from royallepage.rank_cities import DATA_DIR, PROVINCES
from royallepage.scrape import csv_path

COLUMNS = [
    ("full_name", "Full Name"),
    ("phone", "Phone Number"),
    ("rlp_profile_url", "Royal LePage Profile URL"),
    ("personal_website", "Personal/Team Website"),
    ("office", "Office / Brokerage"),
    ("city", "City"),
    ("province", "Province"),
]

OUTPUT_PATH = DATA_DIR / "royal_lepage_agents.xlsx"


def load_rows(prov: str) -> list[dict]:
    path = csv_path(prov)
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def _write_sheet(ws, rows: list[dict]) -> None:
    header_font = Font(bold=True)
    for col_idx, (_, header) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font

    for row_idx, row in enumerate(rows, 2):
        for col_idx, (key, _) in enumerate(COLUMNS, 1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))

    widths = [len(header) for _, header in COLUMNS]
    for row in rows:
        for col_idx, (key, _) in enumerate(COLUMNS):
            widths[col_idx] = max(widths[col_idx], min(len(str(row.get(key, ""))), 60))
    for col_idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width + 2

    ws.freeze_panes = "A2"


def build(output_path: Path = OUTPUT_PATH) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    all_rows: list[dict] = []
    per_province: dict[str, list[dict]] = {}
    for prov, name in PROVINCES.items():
        rows = load_rows(prov)
        rows.sort(key=lambda r: (r.get("city", ""), r.get("full_name", "")))
        per_province[prov] = rows
        all_rows.extend(rows)
        ws = wb.create_sheet(title=name)
        _write_sheet(ws, rows)

    ws_all = wb.create_sheet(title="All", index=0)
    _write_sheet(ws_all, all_rows)

    ws_summary = wb.create_sheet(title="Summary")
    ws_summary.cell(row=1, column=1, value="Province").font = Font(bold=True)
    ws_summary.cell(row=1, column=2, value="City").font = Font(bold=True)
    ws_summary.cell(row=1, column=3, value="Agent Count").font = Font(bold=True)
    row_idx = 2
    for prov, name in PROVINCES.items():
        counts = Counter(r["city"] for r in per_province[prov])
        for city, count in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0])):
            ws_summary.cell(row=row_idx, column=1, value=name)
            ws_summary.cell(row=row_idx, column=2, value=city)
            ws_summary.cell(row=row_idx, column=3, value=count)
            row_idx += 1
        ws_summary.cell(row=row_idx, column=1, value=f"{name} total").font = Font(bold=True)
        ws_summary.cell(row=row_idx, column=3, value=sum(counts.values())).font = Font(bold=True)
        row_idx += 1
    ws_summary.cell(row=row_idx, column=1, value="Grand total").font = Font(bold=True)
    ws_summary.cell(row=row_idx, column=3, value=len(all_rows)).font = Font(bold=True)
    for col_idx, width in enumerate([14, 22, 12], 1):
        ws_summary.column_dimensions[get_column_letter(col_idx)].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def main() -> None:
    path = build()
    print(f"Wrote {path}")


if __name__ == "__main__":
    sys.exit(main())
