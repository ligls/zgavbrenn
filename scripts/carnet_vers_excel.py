#!/usr/bin/env python3
"""Convertit les fiches du carnet en ligne en classeur Excel.

Prend un ou plusieurs dossiers de fiches JSON exportées du magasin du carnet
(un fichier par fiche, nommé d'après l'identifiant PagesJaunes) et produit un
.xlsx à quatre feuilles : les leads, le journal des appels notés, un sommaire
calculé par formules, et un guide des colonnes.

    python scripts/carnet_vers_excel.py --fiches export/leads \
        --sortie leads-carnet-quebec.xlsx

Le carnet en ligne reste la source vivante : ce classeur est une photo. Les
colonnes que tu remplis à la main (statut, rappel, contact, courriel) y sont
reprises telles quelles, pas recalculées.
"""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Le vocabulaire du carnet, pas celui de pj/db.py : les statuts écrits dans le
# magasin viennent de la page du carnet, qui a sa propre liste plus courte.
STATUTS = {
    "a_appeler": "À appeler",
    "rappeler": "À rappeler",
    "injoignable": "Injoignable",
    "interesse": "Intéressé",
    "refus": "Pas intéressé",
    "client": "Client",
}

RESULTATS = {
    "": "—",
    "repondu": "A répondu",
    "vocale": "Boîte vocale",
    "sans_reponse": "Pas de réponse",
    "mauvais_numero": "Mauvais numéro",
    "refus": "Refus",
}

# Les colonnes de la feuille Leads, dans l'ordre : (largeur, libellé).
COLONNES = [
    (38, "Entreprise"), (15, "Téléphone"), (22, "Autres numéros"),
    (34, "Métier"), (30, "Adresse"), (20, "Ville"), (9, "Province"),
    (11, "Code postal"), (34, "Site web"), (15, "Statut"),
    (12, "Rappel le"), (20, "Personne contact"), (24, "Courriel"),
    (12, "Appels notés"), (44, "Dernière note"), (44, "Fiche PagesJaunes"),
]
COL_STATUT = 10
COL_METIER = 4
COL_VILLE = 6
COL_APPELS = 14

POLICE = "Arial"
ENCRE = "1F3B4D"          # bleu ardoise des en-têtes
TRAVAILLE = "FFF4D6"      # fiches où un statut a été changé
TRAIT = Side(style="thin", color="D4D9DE")


def charger(dossiers: list[Path]) -> list[tuple[str, dict]]:
    """Toutes les fiches des dossiers donnés, dédoublonnées par identifiant."""
    fiches: dict[str, dict] = {}
    for dossier in dossiers:
        for chemin in sorted(dossier.glob("*.json")):
            if chemin.name == "lots.json":
                continue
            try:
                fiches[chemin.stem] = json.loads(chemin.read_text(encoding="utf-8"))
            except ValueError as exc:
                raise SystemExit(f"{chemin} : JSON illisible ({exc})")
    return sorted(
        fiches.items(),
        key=lambda paire: (
            (paire[1].get("ville") or "").lower(),
            (paire[1].get("secteur") or "").lower(),
            (paire[1].get("nom") or "").lower(),
        ),
    )


def horodatage(valeur) -> datetime | str:
    """Une date de note ISO en datetime Excel, ou le texte brut si illisible."""
    texte = (valeur or "").strip()
    if not texte:
        return ""
    try:
        return datetime.fromisoformat(texte.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return texte


def jour(valeur) -> datetime | str:
    """Une date de rappel AAAA-MM-JJ en date Excel, ou le texte brut."""
    texte = (valeur or "").strip()
    if not texte:
        return ""
    try:
        return datetime.strptime(texte[:10], "%Y-%m-%d")
    except ValueError:
        return texte


def ligne_lead(fiche: dict) -> list:
    notes = [n for n in (fiche.get("notes") or []) if isinstance(n, dict)]
    tels = [t for t in (fiche.get("tels") or []) if t]
    derniere = ""
    if notes:
        queue = notes[-1]
        derniere = (queue.get("t") or "").strip()
        etiquette = RESULTATS.get(queue.get("r") or "", queue.get("r") or "")
        if etiquette and etiquette != "—":
            derniere = f"{derniere} [{etiquette}]" if derniere else f"[{etiquette}]"
    return [
        fiche.get("nom") or "",
        fiche.get("tel") or "",
        ", ".join(tels[1:]),
        fiche.get("secteur") or "",
        fiche.get("adresse") or "",
        fiche.get("ville") or "",
        fiche.get("province") or "QC",
        fiche.get("cp") or "",
        fiche.get("web") or "",
        STATUTS.get(fiche.get("statut") or "a_appeler", fiche.get("statut") or ""),
        jour(fiche.get("rappel")),
        fiche.get("contact") or "",
        fiche.get("courriel") or "",
        len(notes),
        derniere,
        fiche.get("urlPj") or "",
    ]


def entete(feuille, libelles: list[str]) -> None:
    for index, libelle in enumerate(libelles, start=1):
        cellule = feuille.cell(row=1, column=index, value=libelle)
        cellule.font = Font(name=POLICE, bold=True, color="FFFFFF", size=11)
        cellule.fill = PatternFill("solid", fgColor=ENCRE)
        cellule.alignment = Alignment(vertical="center", horizontal="left")
    feuille.row_dimensions[1].height = 24
    feuille.freeze_panes = "A2"


def feuille_leads(classeur: Workbook, fiches: list[tuple[str, dict]]):
    feuille = classeur.active
    feuille.title = "Leads"
    entete(feuille, [libelle for _, libelle in COLONNES])
    for index, (largeur, _) in enumerate(COLONNES, start=1):
        feuille.column_dimensions[get_column_letter(index)].width = largeur

    for numero, (_, fiche) in enumerate(fiches, start=2):
        valeurs = ligne_lead(fiche)
        travaille = (fiche.get("statut") or "a_appeler") != "a_appeler"
        for colonne, valeur in enumerate(valeurs, start=1):
            cellule = feuille.cell(row=numero, column=colonne, value=valeur)
            cellule.font = Font(name=POLICE, size=10)
            cellule.border = Border(bottom=TRAIT)
            cellule.alignment = Alignment(vertical="top")
            if colonne == COL_STATUT and travaille:
                cellule.font = Font(name=POLICE, size=10, bold=True)
                cellule.fill = PatternFill("solid", fgColor=TRAVAILLE)
            if colonne == COL_APPELS:
                cellule.alignment = Alignment(vertical="top", horizontal="center")
            if colonne == 11 and isinstance(valeur, datetime):
                cellule.number_format = "yyyy-mm-dd"

    derniere = len(fiches) + 1
    feuille.auto_filter.ref = f"A1:{get_column_letter(len(COLONNES))}{derniere}"

    # Le statut se choisit dans une liste : le sommaire compte des libellés
    # exacts, une faute de frappe le ferait mentir.
    liste = DataValidation(
        type="list",
        formula1='"' + ",".join(STATUTS.values()) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    liste.error = "Choisis un statut dans la liste."
    liste.errorTitle = "Statut inconnu"
    feuille.add_data_validation(liste)
    liste.add(f"J2:J{derniere}")
    return feuille


def feuille_appels(classeur: Workbook, fiches: list[tuple[str, dict]]):
    feuille = classeur.create_sheet("Appels")
    colonnes = [
        (22, "Date"), (38, "Entreprise"), (20, "Ville"), (15, "Téléphone"),
        (18, "Résultat"), (15, "Statut actuel"), (70, "Note"),
    ]
    entete(feuille, [libelle for _, libelle in colonnes])
    for index, (largeur, _) in enumerate(colonnes, start=1):
        feuille.column_dimensions[get_column_letter(index)].width = largeur

    rangees = []
    for _, fiche in fiches:
        for note in fiche.get("notes") or []:
            if not isinstance(note, dict):
                continue
            rangees.append((
                horodatage(note.get("d")),
                fiche.get("nom") or "",
                fiche.get("ville") or "",
                fiche.get("tel") or "",
                RESULTATS.get(note.get("r") or "", note.get("r") or ""),
                STATUTS.get(fiche.get("statut") or "a_appeler", ""),
                (note.get("t") or "").strip(),
            ))
    rangees.sort(key=lambda r: (isinstance(r[0], str), str(r[0])), reverse=True)

    for numero, rangee in enumerate(rangees, start=2):
        for colonne, valeur in enumerate(rangee, start=1):
            cellule = feuille.cell(row=numero, column=colonne, value=valeur)
            cellule.font = Font(name=POLICE, size=10)
            cellule.border = Border(bottom=TRAIT)
            cellule.alignment = Alignment(vertical="top", wrap_text=(colonne == 7))
            if colonne == 1 and isinstance(valeur, datetime):
                cellule.number_format = "yyyy-mm-dd hh:mm"

    if rangees:
        feuille.auto_filter.ref = f"A1:G{len(rangees) + 1}"
    else:
        feuille.cell(row=2, column=1, value="Aucun appel noté pour l'instant.").font = (
            Font(name=POLICE, size=10, italic=True)
        )
    return feuille, len(rangees)


def feuille_sommaire(classeur: Workbook, fiches: list[tuple[str, dict]]):
    """Sommaire par formules : il se met à jour quand tu modifies la feuille Leads."""
    feuille = classeur.create_sheet("Sommaire", 0)
    feuille.column_dimensions["A"].width = 44
    feuille.column_dimensions["B"].width = 12
    feuille.sheet_view.showGridLines = False

    fin = len(fiches) + 1
    statut = f"Leads!$J$2:$J${fin}"
    metier = f"Leads!$D$2:$D${fin}"
    ville = f"Leads!$F$2:$F${fin}"

    metiers = sorted({(f.get("secteur") or "") for _, f in fiches})
    compte_villes: dict[str, int] = {}
    for _, f in fiches:
        nom = f.get("ville") or ""
        compte_villes[nom] = compte_villes.get(nom, 0) + 1
    tete = sorted(compte_villes.items(), key=lambda p: (-p[1], p[0].lower()))[:25]

    blocs: list = [
        ("titre", "Carnet de prospection Québec"),
        ("note", f"Photo du carnet en ligne, {datetime.now():%Y-%m-%d}. "
                 "Les totaux sont des formules : ils suivent tes modifications."),
        ("vide", None),
        ("section", "Volume"),
        ("mesure", ("Fiches au total", f"=COUNTA(Leads!$A$2:$A${fin})")),
        # Compte des valeurs distinctes : l'idiome SUMPRODUCT/COUNTIF, qui
        # marche depuis Excel 2007. Le `&""` évite la division par zéro sur
        # les cellules vides.
        ("mesure", ("Municipalités",
                    f'=SUMPRODUCT(({ville}<>"")/COUNTIF({ville},{ville}&""))')),
        ("mesure", ("Métiers",
                    f'=SUMPRODUCT(({metier}<>"")/COUNTIF({metier},{metier}&""))')),
        ("mesure", ("Avec site web", f"=COUNTA(Leads!$I$2:$I${fin})")),
        ("vide", None),
        ("section", "Avancement des appels"),
    ]
    for code, libelle in STATUTS.items():
        blocs.append(("mesure", (libelle, f'=COUNTIF({statut},"{libelle}")')))
    blocs += [
        ("mesure", ("Fiches avec au moins un appel noté",
                    f'=COUNTIF(Leads!$N$2:$N${fin},">0")')),
        ("mesure", ("Appels notés au total", f"=SUM(Leads!$N$2:$N${fin})")),
        ("mesure", ("Rappels planifiés", f"=COUNTA(Leads!$K$2:$K${fin})")),
        ("vide", None),
        ("section", "Par métier"),
    ]
    for nom in metiers:
        blocs.append(("mesure", (nom or "(sans métier)",
                                 f'=COUNTIF({metier},"{nom}")')))
    blocs += [("vide", None), ("section", "25 municipalités les plus fournies")]
    for nom, _ in tete:
        blocs.append(("mesure", (nom, f'=COUNTIF({ville},"{nom}")')))
    reste = "".join(f'-COUNTIF({ville},"{nom}")' for nom, _ in tete)
    blocs.append(("mesure", ("Autres municipalités",
                             f"=COUNTA(Leads!$A$2:$A${fin}){reste}")))

    rangee = 1
    for genre, contenu in blocs:
        if genre == "vide":
            rangee += 1
            continue
        cellule = feuille.cell(row=rangee, column=1)
        if genre == "titre":
            cellule.value = contenu
            cellule.font = Font(name=POLICE, bold=True, size=16, color=ENCRE)
            feuille.row_dimensions[rangee].height = 26
        elif genre == "note":
            cellule.value = contenu
            cellule.font = Font(name=POLICE, size=9, italic=True, color="60686E")
        elif genre == "section":
            cellule.value = contenu
            cellule.font = Font(name=POLICE, bold=True, size=11, color="FFFFFF")
            cellule.fill = PatternFill("solid", fgColor=ENCRE)
            feuille.cell(row=rangee, column=2).fill = PatternFill("solid", fgColor=ENCRE)
        else:
            libelle, formule = contenu
            cellule.value = libelle
            cellule.font = Font(name=POLICE, size=10)
            cellule.border = Border(bottom=TRAIT)
            valeur = feuille.cell(row=rangee, column=2, value=formule)
            valeur.font = Font(name=POLICE, size=10, bold=True)
            valeur.alignment = Alignment(horizontal="right")
            valeur.border = Border(bottom=TRAIT)
        rangee += 1
    return feuille


def feuille_guide(classeur: Workbook, nb_fiches: int, nb_appels: int):
    feuille = classeur.create_sheet("Guide")
    feuille.column_dimensions["A"].width = 26
    feuille.column_dimensions["B"].width = 96
    feuille.sheet_view.showGridLines = False

    lignes: list[tuple[str, str, str]] = [
        ("titre", "Comment lire ce classeur", ""),
        ("texte", "", f"{nb_fiches} fiches, {nb_appels} appel(s) noté(s). "
                      "Le carnet en ligne reste la source vivante ; ce fichier "
                      "est une photo prise le "
                      f"{datetime.now():%Y-%m-%d}. Ce que tu écris ici ne "
                      "remonte pas dans le carnet, et l'inverse non plus."),
        ("vide", "", ""),
        ("section", "Feuilles", ""),
        ("paire", "Sommaire", "Les totaux, calculés par formules sur la feuille "
                              "Leads : ils suivent tes modifications."),
        ("paire", "Leads", "Une ligne par entreprise. Filtres actifs sur la "
                           "rangée d'en-tête."),
        ("paire", "Appels", "Une ligne par appel noté, du plus récent au plus ancien."),
        ("vide", "", ""),
        ("section", "Colonnes relevées sur PagesJaunes", ""),
        ("paire", "Entreprise → Site web",
                  "Extraites de PagesJaunes.ca. Le téléphone est le numéro "
                  "affiché en premier ; sur les fiches payantes c'est parfois un "
                  "numéro de suivi d'appel, le vrai numéro est alors dans "
                  "« Autres numéros »."),
        ("paire", "Fiche PagesJaunes", "Le lien vers la fiche d'origine, pour vérifier."),
        ("vide", "", ""),
        ("section", "Colonnes que tu remplis", ""),
        ("paire", "Statut", "Liste déroulante : "
                            + ", ".join(STATUTS.values())
                            + ". Les fiches déjà travaillées sont surlignées."),
        ("paire", "Rappel le", "Date du prochain appel, format AAAA-MM-JJ."),
        ("paire", "Personne contact", "Qui répond, pour redemander la bonne personne."),
        ("paire", "Courriel", "L'adresse recueillie pendant l'appel."),
        ("paire", "Appels notés / Dernière note",
                  "Recopiés du carnet. Si tu notes tes appels ici plutôt que dans "
                  "le carnet, tiens le compteur à jour : le sommaire s'en sert."),
        ("vide", "", ""),
        ("section", "Résultats d'appel du carnet", ""),
        ("paire", "Vocabulaire",
                  ", ".join(v for k, v in RESULTATS.items() if k)
                  + ". Il apparaît entre crochets dans « Dernière note » et en "
                    "clair dans la feuille Appels."),
        ("vide", "", ""),
        ("section", "Avant d'appeler", ""),
        ("texte", "", "Ce sont des coordonnées d'entreprises publiées "
                      "publiquement, pour de la prospection B2B. Respecte la Loi "
                      "canadienne anti-pourriel (LCAP) et la Liste nationale de "
                      "numéros de télécommunication exclus (LNNTE)."),
    ]

    rangee = 1
    for genre, gauche, droite in lignes:
        if genre == "vide":
            rangee += 1
            continue
        if genre == "titre":
            cellule = feuille.cell(row=rangee, column=1, value=gauche)
            cellule.font = Font(name=POLICE, bold=True, size=16, color=ENCRE)
            feuille.row_dimensions[rangee].height = 26
        elif genre == "section":
            cellule = feuille.cell(row=rangee, column=1, value=gauche)
            cellule.font = Font(name=POLICE, bold=True, size=11, color="FFFFFF")
            cellule.fill = PatternFill("solid", fgColor=ENCRE)
            feuille.cell(row=rangee, column=2).fill = PatternFill("solid", fgColor=ENCRE)
        else:
            if gauche:
                cellule = feuille.cell(row=rangee, column=1, value=gauche)
                cellule.font = Font(name=POLICE, size=10, bold=True)
                cellule.alignment = Alignment(vertical="top")
            texte = feuille.cell(row=rangee, column=2, value=droite)
            texte.font = Font(name=POLICE, size=10)
            texte.alignment = Alignment(vertical="top", wrap_text=True)
            feuille.row_dimensions[rangee].height = 15 * (len(droite) // 95 + 1)
        rangee += 1
    return feuille


def main(argv: list[str] | None = None) -> int:
    parseur = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parseur.add_argument("--fiches", type=Path, action="append", required=True,
                         metavar="DOSSIER",
                         help="dossier de fiches JSON du carnet (répétable)")
    parseur.add_argument("--sortie", type=Path,
                         default=Path("leads-carnet-quebec.xlsx"))
    args = parseur.parse_args(argv)

    for dossier in args.fiches:
        if not dossier.is_dir():
            parseur.error(f"dossier introuvable : {dossier}")

    fiches = charger(args.fiches)
    if not fiches:
        parseur.error("aucune fiche trouvée")

    classeur = Workbook()
    feuille_leads(classeur, fiches)
    _, nb_appels = feuille_appels(classeur, fiches)
    feuille_sommaire(classeur, fiches)
    feuille_guide(classeur, len(fiches), nb_appels)
    classeur.active = 0

    args.sortie.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(args.sortie)

    villes = {(f.get("ville") or "") for _, f in fiches}
    metiers = {(f.get("secteur") or "") for _, f in fiches}
    print(f"{len(fiches)} fiches, {len(villes)} municipalités, {len(metiers)} métiers, "
          f"{nb_appels} appel(s) noté(s) -> {args.sortie}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
