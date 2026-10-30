#!/usr/bin/env python3
"""
Build a collaboration workbook (.xlsx) ready to upload to Google Drive / Sheets.

    python3 importer/export_sheets.py [--out data/Closed_Lost_Q4_2024_collaboration.xlsx]
                                      [--edits path/to/closed_lost_edits.json]

Sheet "Opportunities": one row per opportunity with the work columns first
(Status with a dropdown, Assigned to, Notes, Follow-up), then every source and
derived column.  Header frozen, auto-filter on, Unknown counts highlighted,
status colours.  Sheet "Dashboard": live COUNTIF formulas by bucket, type,
status and close month.  Sheet "Read me": how to use it.

Edits saved from the single-file app (the "N edits" menu -> JSON) or already
stored in data/closed_lost.sqlite are pre-filled when present.
"""

from __future__ import annotations

import argparse
import json
import sqlite3
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
PROJECT_DIR = HERE.parent
DATA_DIR = PROJECT_DIR / "data"

STATUSES = ["Not started", "Called", "Left voicemail", "Callback scheduled", "Not interested", "Won back"]
LINE_TYPES = ["MP", "WLS", "Both", "Unknown"]
BUCKETS = [str(n) for n in range(1, 21)] + ["21-25", "26-30", "31+", "Unknown"]

# (header, key or callable, width)
COLUMNS = [
    ("Account Name", "account_name", 36),
    ("Status", "status", 18),
    ("Assigned to", "assigned_to", 16),
    ("Notes", "notes", 50),
    ("Follow-up", "follow_up_date", 12),
    ("Opportunity Owner", "owner", 22),
    ("Phone", "phone", 16),
    ("Opportunity Name", "opportunity_name", 40),
    ("Close Date", "close_date", 12),
    ("Line Count", "line_count", 11),
    ("Line Type", "line_type", 10),
    ("Bucket", "line_bucket", 9),
    ("Est. Renewal", "estimated_renewal_month_label", 12),
    ("Contract End Hint", "contract_end_hint", 36),
    ("Contract End Month", "contract_end_month", 12),
    ("Additional Info", "additional_info", 60),
    ("Next Step", "next_step", 12),
    ("MP", "mp_count", 6),
    ("WLS", "wls_count", 6),
    ("Count Source", "extraction_source", 14),
    ("Matched Text", "line_match", 20),
    ("Close Month", "close_month_label", 11),
    ("Account Name (raw)", "account_name_raw", 30),
    ("Row Hash", "id", 18),
]
COL = {header: idx + 1 for idx, (header, _, _) in enumerate(COLUMNS)}  # 1-based


def load_edits(sqlite_path: Path, json_path: Path | None) -> dict[str, dict]:
    edits: dict[str, dict] = {}
    if sqlite_path.exists():
        con = sqlite3.connect(sqlite_path)
        con.row_factory = sqlite3.Row
        try:
            for r in con.execute("SELECT * FROM edits"):
                edits[r["row_hash"]] = dict(r)
        except sqlite3.OperationalError:
            pass
        finally:
            con.close()
    if json_path:
        payload = json.loads(json_path.read_text(encoding="utf-8"))
        items = payload.get("edits", payload) if isinstance(payload, dict) else payload
        if isinstance(items, dict):
            items = list(items.values())
        for e in items:
            prev = edits.get(e["row_hash"])
            if not prev or (e.get("updated_at") or "") >= (prev.get("updated_at") or ""):
                edits[e["row_hash"]] = e
    return edits


def to_date(iso: str | None) -> date | None:
    if not iso:
        return None
    try:
        return datetime.strptime(iso[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def build(opps: list[dict], meta: dict, edits: dict[str, dict], out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Opportunities"

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2A78D6")
    work_fill = PatternFill("solid", fgColor="1C5CAB")
    wrap = Alignment(wrap_text=True, vertical="top")

    ws.append([h for h, _, _ in COLUMNS])
    for idx, (header, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = head_font
        cell.fill = work_fill if header in ("Status", "Assigned to", "Notes", "Follow-up") else head_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 22

    for o in opps:
        e = edits.get(o["id"], {})
        count_override = e.get("line_count_override")
        row: list = []
        for header, key, _ in COLUMNS:
            if key == "status":
                row.append(e.get("status") or "Not started")
            elif key == "assigned_to":
                row.append(None)
            elif key == "notes":
                row.append(e.get("notes") or None)
            elif key == "follow_up_date":
                row.append(to_date(e.get("follow_up_date")))
            elif key == "close_date":
                row.append(to_date(o["close_date"]))
            elif key == "line_count":
                row.append(count_override if count_override is not None else o["line_count"])
            elif key == "line_type":
                row.append(e.get("line_type_override") or o["line_type"])
            else:
                row.append(o.get(key))
        ws.append(row)

    n = len(opps) + 1
    last_col = get_column_letter(len(COLUMNS))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{last_col}{n}"

    # date formats + wrapping for long text
    for r in range(2, n + 1):
        ws.cell(row=r, column=COL["Close Date"]).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=COL["Follow-up"]).number_format = "yyyy-mm-dd"
    for header in ("Notes", "Additional Info", "Contract End Hint"):
        for r in range(2, n + 1):
            ws.cell(row=r, column=COL[header]).alignment = wrap

    # dropdowns
    status_dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    status_dv.error = "Choose a status from the list"
    status_dv.prompt = "Work status"
    ws.add_data_validation(status_dv)
    status_dv.add(f"{get_column_letter(COL['Status'])}2:{get_column_letter(COL['Status'])}{n}")
    type_dv = DataValidation(type="list", formula1='"' + ",".join(LINE_TYPES) + '"', allow_blank=True)
    ws.add_data_validation(type_dv)
    type_dv.add(f"{get_column_letter(COL['Line Type'])}2:{get_column_letter(COL['Line Type'])}{n}")
    date_dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True)
    ws.add_data_validation(date_dv)
    date_dv.add(f"{get_column_letter(COL['Follow-up'])}2:{get_column_letter(COL['Follow-up'])}{n}")

    # conditional formatting
    st = get_column_letter(COL["Status"])
    lc = get_column_letter(COL["Line Count"])
    status_rng = f"{st}2:{st}{n}"
    ws.conditional_formatting.add(status_rng, CellIsRule(operator="equal", formula=['"Won back"'], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(color="006100")))
    ws.conditional_formatting.add(status_rng, CellIsRule(operator="equal", formula=['"Not interested"'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(color="9C0006")))
    ws.conditional_formatting.add(status_rng, CellIsRule(operator="equal", formula=['"Callback scheduled"'], fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(color="7A5300")))
    ws.conditional_formatting.add(status_rng, CellIsRule(operator="equal", formula=['"Called"'], fill=PatternFill("solid", fgColor="CDE2FB")))
    ws.conditional_formatting.add(f"{lc}2:{lc}{n}", FormulaRule(formula=[f"ISBLANK({lc}2)"], fill=PatternFill("solid", fgColor="FFEB9C")))
    fu = get_column_letter(COL["Follow-up"])
    ws.conditional_formatting.add(f"{fu}2:{fu}{n}", FormulaRule(formula=[f"AND(NOT(ISBLANK({fu}2)),{fu}2<=TODAY())"], fill=PatternFill("solid", fgColor="FFC7CE")))

    # ------------------------------------------------------------ dashboard
    dash = wb.create_sheet("Dashboard")
    dash.column_dimensions["A"].width = 26
    dash.column_dimensions["B"].width = 12
    dash.column_dimensions["D"].width = 26
    dash.column_dimensions["E"].width = 12
    dash.column_dimensions["G"].width = 26
    dash.column_dimensions["H"].width = 12
    bold = Font(bold=True)
    data = "Opportunities"
    lc_rng = f"{data}!${lc}$2:${lc}${n}"
    st_rng = f"{data}!${st}$2:${st}${n}"
    lt = get_column_letter(COL["Line Type"])
    lt_rng = f"{data}!${lt}$2:${lt}${n}"
    cm = get_column_letter(COL["Close Month"])
    cm_rng = f"{data}!${cm}$2:${cm}${n}"
    fu_rng = f"{data}!${fu}$2:${fu}${n}"

    dash["A1"], dash["B1"] = "Summary", None
    dash["A1"].font = bold
    dash["A2"], dash["B2"] = "Opportunities", f"=COUNTA({data}!$A$2:$A${n})"
    dash["A3"], dash["B3"] = "Unknown line count", f"=COUNTBLANK({lc_rng})"
    dash["A4"], dash["B4"] = "Known lines (sum)", f"=SUM({lc_rng})"
    dash["A5"], dash["B5"] = "Worked (status set)", f'=COUNTIFS({st_rng},"<>Not started")'
    dash["A6"], dash["B6"] = "Follow-ups due", f'=COUNTIFS({fu_rng},"<="&TODAY())'
    dash["A7"], dash["B7"] = "Won back", f'=COUNTIF({st_rng},"Won back")'

    dash["A9"] = "By bucket"
    dash["A9"].font = bold
    r = 10
    for b in BUCKETS:
        dash.cell(row=r, column=1, value=b)
        if b == "Unknown":
            f = f"=COUNTBLANK({lc_rng})"
        elif b == "31+":
            f = f'=COUNTIF({lc_rng},">=31")'
        elif "-" in b:
            lo, hi = b.split("-")
            f = f'=COUNTIFS({lc_rng},">={lo}",{lc_rng},"<={hi}")'
        else:
            f = f"=COUNTIF({lc_rng},{b})"
        dash.cell(row=r, column=2, value=f)
        r += 1

    dash["D9"] = "By line type"
    dash["D9"].font = bold
    for i, t in enumerate(LINE_TYPES, start=10):
        dash.cell(row=i, column=4, value=t)
        dash.cell(row=i, column=5, value=f'=COUNTIF({lt_rng},"{t}")')
    dash["D16"] = "By status"
    dash["D16"].font = bold
    for i, s in enumerate(STATUSES, start=17):
        dash.cell(row=i, column=4, value=s)
        dash.cell(row=i, column=5, value=f'=COUNTIF({st_rng},"{s}")')

    dash["G9"] = "By close month"
    dash["G9"].font = bold
    months = sorted({o["close_month_label"] for o in opps}, key=lambda m: [o["close_month"] for o in opps if o["close_month_label"] == m][0])
    for i, m in enumerate(months, start=10):
        dash.cell(row=i, column=7, value=m)
        dash.cell(row=i, column=8, value=f'=COUNTIF({cm_rng},"{m}")')
    dash["G16"] = "By owner (top)"
    dash["G16"].font = bold
    ow = get_column_letter(COL["Opportunity Owner"])
    dash["G17"] = "Use Data > Pivot table on the Opportunities sheet (Owner x Status) for the full breakdown."
    dash["G17"].alignment = Alignment(wrap_text=True)
    dash.row_dimensions[17].height = 42
    _ = ow

    # ------------------------------------------------------------ read me
    rm = wb.create_sheet("Read me")
    rm.column_dimensions["A"].width = 110
    lines = [
        ("Closed Lost Q4 2024 - renewal workbench (collaboration copy)", True),
        (f"Source: {meta.get('source_file')} - imported {meta.get('imported_at', '')[:16].replace('T', ' ')} - {len(opps):,} opportunities"
         + (f" - {meta.get('excluded_count')} {meta.get('excluded_label') or ''} opportunities excluded" if meta.get("excluded_count") else ""), False),
        ("", False),
        ("How to work in this sheet", True),
        ("1. Upload this file to Google Drive and open it with Google Sheets (File > Save as Google Sheets), then share it with your colleague.", False),
        ("2. The blue columns B-E are yours: Status (dropdown), Assigned to, Notes, Follow-up (date). Everything else comes from Salesforce / the parser.", False),
        ("3. Use the filter arrows on row 1 (or Data > Create a filter view so each person keeps their own filters).", False),
        ("4. Line Count / Line Type were parsed from the Opportunity Name (then Additional Info). Blank Line Count = unknown (highlighted): fix it by typing the number.", False),
        ("5. The Dashboard sheet updates automatically from the Status, Line Count and Follow-up columns.", False),
        ("6. Row Hash identifies a row across re-exports; keep that column if you want to merge changes back into the app later.", False),
        ("", False),
        ("Statuses: " + " | ".join(STATUSES), False),
        ("Buckets: 1 ... 20, 21-25, 26-30, 31+, Unknown. Est. Renewal = close date + 24 months.", False),
    ]
    for i, (text, is_bold) in enumerate(lines, start=1):
        c = rm.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if is_bold:
            c.font = Font(bold=True, size=12 if i == 1 else 11)

    wb.move_sheet("Read me", offset=-2)
    wb.active = 1
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data", default=str(DATA_DIR / "opportunities.json"))
    ap.add_argument("--db", default=str(DATA_DIR / "closed_lost.sqlite"))
    ap.add_argument("--edits", default=None, help="edits JSON saved from the single-file app (merged with SQLite edits, newest wins)")
    ap.add_argument("--out", default=str(DATA_DIR / "Closed_Lost_Q4_2024_collaboration.xlsx"))
    args = ap.parse_args()

    data_path = Path(args.data)
    if not data_path.exists():
        sys.exit(f"{data_path} not found - run the importer first (make import).")
    payload = json.loads(data_path.read_text(encoding="utf-8"))
    edits = load_edits(Path(args.db), Path(args.edits) if args.edits else None)
    out = Path(args.out)
    build(payload["opportunities"], payload["meta"], edits, out)
    print(f"Wrote {out} ({out.stat().st_size / 1e6:.1f} MB): {len(payload['opportunities']):,} rows, {len(edits)} pre-filled edit(s).")
    print("Upload it to Google Drive, open with Google Sheets, share with your colleague.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
