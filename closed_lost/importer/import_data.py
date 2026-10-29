#!/usr/bin/env python3
"""
Import the Salesforce "Closed Lost Q4 2024" export into the local data store.

    python3 importer/import_data.py [path/to/Closed_Lost_*.xlsx] [--out-dir data]

What it does
------------
1. Reads the single sheet, finds the header row ("Opportunity Owner" ...),
   skips the report metadata above it and the "Total"/footer rows below it.
2. Cleans every row: whitespace, mojibake repair of Account/Opportunity
   names (raw values are kept), phone -> tel: link, close date -> ISO.
3. Derives: line_count / line_type / bucket / extraction source / matched
   snippet (``line_parser``), close month, estimated renewal (+24 months),
   contract-end hint (``contract_hints``), and a stable ``row_hash``
   (owner + account + opportunity name + close date) used to key edits.
4. Writes ``data/opportunities.json`` (what the web app loads) and refreshes
   the ``opportunities`` table of ``data/closed_lost.sqlite`` - the ``edits``
   table in the same database is never touched, so statuses, notes and
   overrides survive re-imports.
5. Optionally drops rows whose Opportunity Name matches ``--exclude-name``
   (case-insensitive regex), e.g. ``--exclude-name renewal``.
6. Prints (and saves to ``data/import_report.txt``) a validation report:
   row count, extraction sources, bucket distribution and 20 random rows
   with the original text next to the parsed count/type.
"""

from __future__ import annotations

import argparse
import collections
import glob
import hashlib
import json
import os
import random
import re
import sqlite3
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from contract_hints import extract_contract_end  # noqa: E402
from line_parser import BUCKETS, parse_opportunity  # noqa: E402
from text_repair import clean_whitespace, repair_mojibake  # noqa: E402

PARSER_VERSION = "1.0"
EXPECTED_HEADERS = [
    "Opportunity Owner", "Account Name", "Phone", "Opportunity Name",
    "Close Date", "Next Step", "Additional Info",
]
FOOTER_MARKERS = ("Confidential Information", "Copyright ©", "salesforce.com")
MONTH_ABBR = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

HERE = Path(__file__).resolve().parent
PROJECT_DIR = HERE.parent          # closed_lost/
REPO_DIR = PROJECT_DIR.parent      # repository root


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------

def find_source_file(explicit: str | None) -> Path:
    if explicit:
        p = Path(explicit)
        if not p.exists():
            sys.exit(f"Source file not found: {p}")
        return p
    candidates: list[str] = []
    for folder in (Path.cwd(), REPO_DIR, PROJECT_DIR, PROJECT_DIR / "data"):
        candidates += glob.glob(str(folder / "Closed_Lost*.xlsx"))
    candidates = sorted(set(candidates), key=lambda f: os.path.getmtime(f), reverse=True)
    if not candidates:
        sys.exit("No Closed_Lost*.xlsx found. Pass the path explicitly: "
                 "python3 importer/import_data.py path/to/file.xlsx")
    return Path(candidates[0])


def parse_close_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def add_months(d: date, months: int) -> date:
    month_index = d.month - 1 + months
    year = d.year + month_index // 12
    month = month_index % 12 + 1
    # clamp the day to the target month's length
    for day in (d.day, 30, 29, 28):
        try:
            return date(year, month, day)
        except ValueError:
            continue
    raise AssertionError("unreachable")


def month_label(d: date) -> str:
    return f"{MONTH_ABBR[d.month - 1]} {d.year}"


def month_key(d: date) -> str:
    return f"{d.year:04d}-{d.month:02d}"


def phone_to_tel(phone: str | None) -> str | None:
    if not phone:
        return None
    digits = re.sub(r"\D", "", phone)
    if not digits:
        return None
    if len(digits) == 10:
        return f"+1{digits}"
    if len(digits) == 11 and digits.startswith("1"):
        return f"+{digits}"
    return digits


def stable_hash(*parts: str) -> str:
    norm = "\x1f".join(re.sub(r"\s+", " ", (p or "")).strip().casefold() for p in parts)
    return hashlib.sha1(norm.encode("utf-8")).hexdigest()[:16]


def is_footer(cells: list) -> bool:
    joined = " ".join(str(c) for c in cells if c is not None)
    return any(marker in joined for marker in FOOTER_MARKERS)


# --------------------------------------------------------------------------
# Reading the workbook
# --------------------------------------------------------------------------

def read_rows(path: Path) -> tuple[list[dict], dict]:
    """Return (records, stats) from the xlsx export."""
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")      # openpyxl: "no default style"
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_row_no = None
    col_index: dict[str, int] = {}
    records: list[dict] = []
    stats = collections.Counter()

    for row_no, cells in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = list(cells)
        if header_row_no is None:
            texts = [str(c).strip() if c is not None else "" for c in cells]
            if "Opportunity Owner" in texts:
                header_row_no = row_no
                for idx, text in enumerate(texts):
                    if text:
                        col_index[text] = idx
                missing = [h for h in EXPECTED_HEADERS if h not in col_index]
                if missing:
                    sys.exit(f"Header row {row_no} is missing columns: {missing}")
            continue

        def cell(name: str):
            idx = col_index.get(name)
            return cells[idx] if idx is not None and idx < len(cells) else None

        owner = clean_whitespace(cell("Opportunity Owner"))
        if not owner:
            stats["skipped_blank_owner"] += 1
            continue
        if is_footer(cells):
            stats["skipped_footer"] += 1
            continue
        close = parse_close_date(cell("Close Date"))
        account_raw = cell("Account Name")
        if close is None or not isinstance(account_raw, str):
            # the "Total | Count | 5490" summary row lands here
            stats["skipped_summary"] += 1
            continue

        records.append({
            "row_number": row_no,
            "owner": owner,
            "account_raw": clean_whitespace(account_raw) or "",
            "phone": clean_whitespace(cell("Phone")),
            "opportunity_raw": clean_whitespace(cell("Opportunity Name")) or "",
            "close": close,
            "next_step": clean_whitespace(cell("Next Step")),
            "additional_info": clean_whitespace(cell("Additional Info")),
        })

    if header_row_no is None:
        sys.exit("Could not find the header row (a row containing 'Opportunity Owner').")
    stats["header_row"] = header_row_no
    wb.close()
    return records, dict(stats)


# --------------------------------------------------------------------------
# Deriving fields
# --------------------------------------------------------------------------

def build_opportunities(raw_records: list[dict]) -> tuple[list[dict], dict]:
    stats = collections.Counter()
    seen_hashes: collections.Counter = collections.Counter()
    opportunities: list[dict] = []

    for rec in raw_records:
        account, acc_fixed = repair_mojibake(rec["account_raw"])
        opp_name, opp_fixed = repair_mojibake(rec["opportunity_raw"])
        info, _ = repair_mojibake(rec["additional_info"])
        stats["repaired_account_names"] += int(acc_fixed)
        stats["repaired_opportunity_names"] += int(opp_fixed)

        close: date = rec["close"]
        base_key = stable_hash(rec["owner"], rec["account_raw"], rec["opportunity_raw"], close.isoformat())
        seen_hashes[base_key] += 1
        row_hash = base_key if seen_hashes[base_key] == 1 else stable_hash(base_key, f"#{seen_hashes[base_key]}")
        if seen_hashes[base_key] > 1:
            stats["duplicate_rows_disambiguated"] += 1

        lines = parse_opportunity(opp_name, info)
        hint = extract_contract_end(info)
        renewal = add_months(close, 24)

        opportunities.append({
            "id": row_hash,
            "row_number": rec["row_number"],
            "owner": rec["owner"],
            "account_name": account,
            "account_name_raw": rec["account_raw"],
            "account_name_repaired": acc_fixed,
            "phone": rec["phone"],
            "phone_tel": phone_to_tel(rec["phone"]),
            "opportunity_name": opp_name,
            "opportunity_name_raw": rec["opportunity_raw"],
            "close_date": close.isoformat(),
            "close_month": month_key(close),
            "close_month_label": month_label(close),
            "next_step": rec["next_step"],
            "additional_info": info,
            "line_count": lines.line_count,
            "line_type": lines.line_type,
            "mp_count": lines.mp_count,
            "wls_count": lines.wls_count,
            "types_mentioned": lines.types_mentioned,
            "extraction_source": lines.extraction_source,
            "line_match": lines.line_match,
            "line_bucket": lines.line_bucket,
            "estimated_renewal_date": renewal.isoformat(),
            "estimated_renewal_month": month_key(renewal),
            "estimated_renewal_month_label": month_label(renewal),
            "contract_end_hint": hint.hint if hint else None,
            "contract_end_month": hint.month if hint else None,
            "contract_end_confidence": hint.confidence if hint else None,
        })
    return opportunities, dict(stats)


# --------------------------------------------------------------------------
# Persistence
# --------------------------------------------------------------------------

SQL_SCHEMA = """
CREATE TABLE IF NOT EXISTS edits (
    row_hash            TEXT PRIMARY KEY,
    line_count_override INTEGER,
    line_type_override  TEXT,
    status              TEXT NOT NULL DEFAULT 'Not started',
    notes               TEXT NOT NULL DEFAULT '',
    follow_up_date      TEXT,
    updated_at          TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS imports (
    imported_at  TEXT NOT NULL,
    source_file  TEXT NOT NULL,
    row_count    INTEGER NOT NULL,
    parser_version TEXT NOT NULL
);
"""

OPP_COLUMNS = [
    "id", "row_number", "owner", "account_name", "account_name_raw", "account_name_repaired",
    "phone", "phone_tel", "opportunity_name", "opportunity_name_raw", "close_date", "close_month",
    "close_month_label", "next_step", "additional_info", "line_count", "line_type", "mp_count",
    "wls_count", "types_mentioned", "extraction_source", "line_match", "line_bucket",
    "estimated_renewal_date", "estimated_renewal_month", "estimated_renewal_month_label",
    "contract_end_hint", "contract_end_month", "contract_end_confidence",
]


def write_sqlite(db_path: Path, opportunities: list[dict], meta: dict) -> int:
    """Replace the opportunities table, keep edits. Returns # of orphaned edits."""
    con = sqlite3.connect(db_path)
    try:
        con.executescript(SQL_SCHEMA)
        con.execute("DROP TABLE IF EXISTS opportunities")
        cols_sql = ", ".join(
            f"{c} {'INTEGER' if c in ('row_number', 'line_count', 'mp_count', 'wls_count') else 'TEXT'}"
            + (" PRIMARY KEY" if c == "id" else "")
            for c in OPP_COLUMNS
        )
        con.execute(f"CREATE TABLE opportunities ({cols_sql})")
        rows = []
        for o in opportunities:
            row = []
            for c in OPP_COLUMNS:
                v = o[c]
                if isinstance(v, list):
                    v = ",".join(v)
                elif isinstance(v, bool):
                    v = int(v)
                row.append(v)
            rows.append(row)
        con.executemany(
            f"INSERT INTO opportunities ({', '.join(OPP_COLUMNS)}) VALUES ({', '.join('?' * len(OPP_COLUMNS))})",
            rows,
        )
        con.execute("CREATE INDEX idx_opp_owner ON opportunities(owner)")
        con.execute("CREATE INDEX idx_opp_bucket ON opportunities(line_bucket)")
        con.execute(
            "INSERT INTO imports (imported_at, source_file, row_count, parser_version) VALUES (?, ?, ?, ?)",
            (meta["imported_at"], meta["source_file"], meta["row_count"], PARSER_VERSION),
        )
        orphans = con.execute(
            "SELECT COUNT(*) FROM edits WHERE row_hash NOT IN (SELECT id FROM opportunities)"
        ).fetchone()[0]
        con.commit()
        return orphans
    finally:
        con.close()


# --------------------------------------------------------------------------
# Validation report
# --------------------------------------------------------------------------

def shorten(text: str | None, width: int) -> str:
    if not text:
        return ""
    text = re.sub(r"\s+", " ", text)
    return text if len(text) <= width else text[: width - 1] + "…"


def build_report(opportunities: list[dict], read_stats: dict, build_stats: dict, meta: dict,
                 orphans: int, seed: int, sample_size: int) -> str:
    out: list[str] = []
    n = len(opportunities)
    out.append("=" * 100)
    out.append("IMPORT VALIDATION REPORT")
    out.append("=" * 100)
    out.append(f"Source file        : {meta['source_file']}")
    out.append(f"Header row         : {read_stats.get('header_row')}")
    out.append(f"Opportunities      : {n:,}")
    if meta.get("excluded_name_pattern"):
        out.append(f"Excluded by name   : {meta['excluded_count']:,} rows whose Opportunity Name matches /{meta['excluded_name_pattern']}/i")
    out.append(f"Skipped rows       : blank owner={read_stats.get('skipped_blank_owner', 0)}, "
               f"footer={read_stats.get('skipped_footer', 0)}, summary/total={read_stats.get('skipped_summary', 0)}")
    out.append(f"Repaired names     : accounts={build_stats.get('repaired_account_names', 0)}, "
               f"opportunity names={build_stats.get('repaired_opportunity_names', 0)}")
    out.append(f"Duplicate rows     : {build_stats.get('duplicate_rows_disambiguated', 0)} (same owner/account/opp/date; hashed with a suffix)")
    out.append(f"Owners             : {len({o['owner'] for o in opportunities}):,}")
    out.append(f"With phone         : {sum(1 for o in opportunities if o['phone']):,}")
    out.append(f"With Additional Info: {sum(1 for o in opportunities if o['additional_info']):,}")
    out.append(f"With contract hint : {sum(1 for o in opportunities if o['contract_end_hint']):,} "
               f"(high={sum(1 for o in opportunities if o['contract_end_confidence'] == 'high'):,}, "
               f"low={sum(1 for o in opportunities if o['contract_end_confidence'] == 'low'):,}, "
               f"with month={sum(1 for o in opportunities if o['contract_end_month']):,})")
    out.append(f"Orphaned edits     : {orphans} (edits whose row no longer exists in this import)")
    out.append("")

    out.append("Line count extraction source")
    src = collections.Counter(o["extraction_source"] for o in opportunities)
    for key, label in (("name", "from Opportunity Name"), ("additional_info", "from Additional Info"), ("none", "none (Unknown count)")):
        c = src.get(key, 0)
        out.append(f"  {label:<24} {c:>6,}  ({c / n:6.1%})")
    out.append("")

    out.append("Line type")
    types = collections.Counter(o["line_type"] for o in opportunities)
    for key in ("MP", "WLS", "Both", "Unknown"):
        c = types.get(key, 0)
        out.append(f"  {key:<24} {c:>6,}  ({c / n:6.1%})")
    out.append("")

    out.append("Close month")
    months = collections.Counter(o["close_month_label"] for o in opportunities)
    for key in sorted(months, key=lambda k: [o["close_month"] for o in opportunities if o["close_month_label"] == k][0]):
        out.append(f"  {key:<24} {months[key]:>6,}")
    out.append("")

    out.append("Bucket distribution")
    buckets = collections.Counter(o["line_bucket"] for o in opportunities)
    width = max(buckets.values()) if buckets else 1
    for b in BUCKETS:
        c = buckets.get(b, 0)
        bar = "#" * int(round(40 * c / width)) if width else ""
        out.append(f"  {b:<8} {c:>6,}  {bar}")
    total_lines = sum(o["line_count"] for o in opportunities if o["line_count"])
    out.append(f"  Total lines with a known count: {total_lines:,}")
    out.append("")

    rng = random.Random(seed)
    sample = rng.sample(opportunities, min(sample_size, n))
    out.append(f"{sample_size} random rows (seed {seed}) - original text vs parsed result")
    out.append("-" * 100)
    for o in sample:
        out.append(f"row {o['row_number']:>5} | {shorten(o['opportunity_name'], 70)}")
        if o["extraction_source"] == "additional_info" or (o["extraction_source"] == "none" and o["additional_info"]):
            out.append(f"          | info: {shorten(o['additional_info'], 90)}")
        out.append(f"          -> count={o['line_count']!s:<5} type={o['line_type']:<8} bucket={o['line_bucket']:<8} "
                   f"source={o['extraction_source']:<15} matched={o['line_match'] or '-'}")
    out.append("-" * 100)
    return "\n".join(out)


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------

def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", nargs="?", help="Salesforce export (default: newest Closed_Lost*.xlsx in the repo)")
    ap.add_argument("--out-dir", default=str(PROJECT_DIR / "data"), help="output folder (default: closed_lost/data)")
    ap.add_argument("--seed", type=int, default=42, help="seed for the random validation sample")
    ap.add_argument("--sample", type=int, default=20, help="number of random rows in the report")
    ap.add_argument("--exclude-name", metavar="REGEX", default=None,
                    help="drop opportunities whose Opportunity Name matches this case-insensitive regex, e.g. 'renewal'")
    ap.add_argument("--exclude-label", metavar="TEXT", default=None,
                    help="short human label for the exclusion shown in the app, e.g. 'renewal / RNW'")
    args = ap.parse_args(argv)

    source = find_source_file(args.xlsx)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    print(f"Reading {source} ...")
    raw_records, read_stats = read_rows(source)
    opportunities, build_stats = build_opportunities(raw_records)

    excluded = 0
    if args.exclude_name:
        rx = re.compile(args.exclude_name, re.IGNORECASE)
        kept = [o for o in opportunities if not rx.search(o["opportunity_name"] or "")]
        excluded = len(opportunities) - len(kept)
        opportunities = kept

    meta = {
        "source_file": source.name,
        "imported_at": datetime.now().isoformat(timespec="seconds"),
        "row_count": len(opportunities),
        "parser_version": PARSER_VERSION,
        "excluded_name_pattern": args.exclude_name,
        "excluded_count": excluded,
        "excluded_label": args.exclude_label if args.exclude_name else None,
        "buckets": BUCKETS,
        "line_types": ["MP", "WLS", "Both", "Unknown"],
        "statuses": ["Not started", "Called", "Left voicemail", "Callback scheduled", "Not interested", "Won back"],
    }

    json_path = out_dir / "opportunities.json"
    with open(json_path, "w", encoding="utf-8") as fh:
        json.dump({"meta": meta, "opportunities": opportunities}, fh, ensure_ascii=False, separators=(",", ":"))
    db_path = out_dir / "closed_lost.sqlite"
    orphans = write_sqlite(db_path, opportunities, meta)

    report = build_report(opportunities, read_stats, build_stats, meta, orphans, args.seed, args.sample)
    (out_dir / "import_report.txt").write_text(report + "\n", encoding="utf-8")
    print(report)
    print(f"\nWrote {json_path} ({json_path.stat().st_size / 1e6:.1f} MB) and {db_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
